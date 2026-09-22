from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "deliverables" / "Kimmy-Nails_Huong-dan-van-hanh-website-va-booking.docx"

# compact_reference_guide, with a named Kimmy Nails brand-palette override.
FONT = "Calibri"
INK = "26211C"
MUTED = "6D655E"
GOLD = "9A6A2F"
GOLD_LIGHT = "F4EBDD"
CREAM = "FFF9F1"
LINE = "DCCDBA"
WHITE = "FFFFFF"
GREEN = "2F6B4F"
RED = "9B3A32"
BLUE = "315E7A"
TABLE_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120


def rgb(hex_value):
    return RGBColor.from_string(hex_value)


def set_run_font(run, size=None, bold=None, italic=None, color=INK, name=FONT):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color:
        run.font.color.rgb = rgb(color)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=LINE, size=6):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), str(size))
        tag.set(qn("w:space"), "0")
        tag.set(qn("w:color"), color)


def set_table_geometry(table, widths, indent=TABLE_INDENT_DXA):
    assert sum(widths) == TABLE_WIDTH_DXA, (widths, sum(widths))
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(TABLE_WIDTH_DXA))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent))
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        cant_split = OxmlElement("w:cantSplit")
        row._tr.get_or_add_trPr().append(cant_split)
        for cell, width in zip(row.cells, widths):
            cell.width = Inches(width / 1440)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.first_child_found_in("w:tcW")
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell)


def mark_header_row(row):
    tr_pr = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    tr_pr.append(repeat)


def add_numbering_definition(doc, ordered=False, start_at=1):
    numbering = doc.part.numbering_part.element
    abstract_ids = [int(x.get(qn("w:abstractNumId"))) for x in numbering.findall(qn("w:abstractNum"))]
    num_ids = [int(x.get(qn("w:numId"))) for x in numbering.findall(qn("w:num"))]
    abstract_id = max(abstract_ids, default=0) + 1
    num_id = max(num_ids, default=0) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    nsid = OxmlElement("w:nsid")
    nsid.set(qn("w:val"), f"{abstract_id:08X}")
    abstract.append(nsid)
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), str(start_at))
    level.append(start)
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), "decimal" if ordered else "bullet")
    level.append(num_fmt)
    lvl_text = OxmlElement("w:lvlText")
    lvl_text.set(qn("w:val"), "%1." if ordered else "•")
    level.append(lvl_text)
    suffix = OxmlElement("w:suff")
    suffix.set(qn("w:val"), "tab")
    level.append(suffix)
    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "540")
    tabs.append(tab)
    p_pr.append(tabs)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "540")
    ind.set(qn("w:hanging"), "270")
    p_pr.append(ind)
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:after"), "80")
    spacing.set(qn("w:line"), "300")
    spacing.set(qn("w:lineRule"), "auto")
    p_pr.append(spacing)
    level.append(p_pr)
    r_pr = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), FONT)
    fonts.set(qn("w:hAnsi"), FONT)
    r_pr.append(fonts)
    level.append(r_pr)
    abstract.append(level)
    # OOXML requires every abstractNum definition to appear before concrete num
    # instances. Word otherwise ignores the custom definition and continues an
    # unrelated built-in list, which breaks both bullets and restarts.
    first_num_index = next(
        (index for index, child in enumerate(numbering) if child.tag == qn("w:num")),
        len(numbering),
    )
    numbering.insert(first_num_index, abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_num_id = OxmlElement("w:abstractNumId")
    abstract_num_id.set(qn("w:val"), str(abstract_id))
    num.append(abstract_num_id)
    if ordered:
        override = OxmlElement("w:lvlOverride")
        override.set(qn("w:ilvl"), "0")
        start_override = OxmlElement("w:startOverride")
        start_override.set(qn("w:val"), str(start_at))
        override.append(start_override)
        num.append(override)
    numbering.append(num)
    return num_id


def apply_num(paragraph, num_id):
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num = OxmlElement("w:numId")
    num.set(qn("w:val"), str(num_id))
    num_pr.append(ilvl)
    num_pr.append(num)
    p_pr.append(num_pr)
    paragraph.paragraph_format.left_indent = Inches(0.375)
    paragraph.paragraph_format.first_line_indent = Inches(-0.188)
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.line_spacing = 1.25


def add_list_item(doc, text, num_id, bold_prefix=None):
    p = doc.add_paragraph()
    apply_num(p, num_id)
    if bold_prefix and text.startswith(bold_prefix):
        first = p.add_run(bold_prefix)
        set_run_font(first, bold=True)
        rest = p.add_run(text[len(bold_prefix):])
        set_run_font(rest)
    else:
        set_run_font(p.add_run(text))
    return p


def add_para(doc, text="", *, size=11, bold=False, italic=False, color=INK,
             align=None, before=0, after=6, line=1.25, keep_next=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = line
    p.paragraph_format.keep_with_next = keep_next
    if align is not None:
        p.alignment = align
    set_run_font(p.add_run(text), size=size, bold=bold, italic=italic, color=color)
    return p


def add_mixed_para(doc, parts, *, before=0, after=6, line=1.25):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = line
    for text, options in parts:
        set_run_font(p.add_run(text), **options)
    return p


def add_hyperlink(paragraph, text, url, color=BLUE, underline=True):
    relationship_id = paragraph.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    r_fonts = OxmlElement("w:rFonts")
    r_fonts.set(qn("w:ascii"), FONT)
    r_fonts.set(qn("w:hAnsi"), FONT)
    r_pr.append(r_fonts)
    color_node = OxmlElement("w:color")
    color_node.set(qn("w:val"), color)
    r_pr.append(color_node)
    if underline:
        underline_node = OxmlElement("w:u")
        underline_node.set(qn("w:val"), "single")
        r_pr.append(underline_node)
    run.append(r_pr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    set_run_font(p.add_run(text), bold=True, color=GOLD if level < 3 else INK)
    return p


def add_callout(doc, label, body, fill=CREAM, accent=GOLD):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [TABLE_WIDTH_DXA])
    set_table_borders(table, color=accent, size=8)
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.2
    set_run_font(p.add_run(f"{label}: "), bold=True, color=accent)
    set_run_font(p.add_run(body), color=INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_table(doc, headers, rows, widths, alignments=None):
    table = doc.add_table(rows=1, cols=len(headers))
    set_table_geometry(table, widths)
    set_table_borders(table)
    mark_header_row(table.rows[0])
    for index, (cell, value) in enumerate(zip(table.rows[0].cells, headers)):
        set_cell_shading(cell, GOLD_LIGHT)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.1
        set_run_font(p.add_run(value), size=10, bold=True, color=INK)
    for row_values in rows:
        cells = table.add_row().cells
        for index, (cell, value) in enumerate(zip(cells, row_values)):
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.15
            if alignments:
                p.alignment = alignments[index]
            set_run_font(p.add_run(value), size=10, color=INK)
    set_table_geometry(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_run_font(paragraph.add_run("Trang "), size=9, color=MUTED)
    run = paragraph.add_run()
    fld_char_1 = OxmlElement("w:fldChar")
    fld_char_1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char_2 = OxmlElement("w:fldChar")
    fld_char_2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char_1)
    run._r.append(instr_text)
    run._r.append(fld_char_2)
    set_run_font(run, size=9, color=MUTED)


def configure_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    normal.font.size = Pt(11)
    normal.font.color.rgb = rgb(INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    configs = {
        "Heading 1": (16, GOLD, 18, 10),
        "Heading 2": (13, GOLD, 14, 7),
        "Heading 3": (12, INK, 10, 5),
    }
    for name, (size, color, before, after) in configs.items():
        style = styles[name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:ascii"), FONT)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = rgb(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True


def build_document():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    section.different_first_page_header_footer = True
    configure_styles(doc)
    bullet_id = add_numbering_definition(doc, ordered=False)
    cancel_flow_number_ids = [add_numbering_definition(doc, ordered=True, start_at=i) for i in range(1, 7)]
    admin_search_number_ids = [add_numbering_definition(doc, ordered=True, start_at=i) for i in range(1, 6)]
    admin_cancel_number_ids = [add_numbering_definition(doc, ordered=True, start_at=i) for i in range(1, 5)]

    props = doc.core_properties
    props.title = "Hướng dẫn vận hành website và hệ thống booking Kimmy Nails"
    props.subject = "Luồng đặt lịch, WhatsApp, hủy lịch và quản trị booking"
    props.author = "Kimmy Nails"
    props.keywords = "Kimmy Nails, booking, Supabase, Vercel, WhatsApp, admin"

    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_run_font(hp.add_run("KIMMY NAILS  /  HƯỚNG DẪN VẬN HÀNH"), size=8.5, bold=True, color=MUTED)
    fp = section.footer.paragraphs[0]
    add_page_number(fp)

    # Editorial-cover opening.
    add_para(doc, "TÀI LIỆU VẬN HÀNH NỘI BỘ", size=10, bold=True, color=GOLD,
             align=WD_ALIGN_PARAGRAPH.CENTER, before=42, after=18)
    add_para(doc, "Hướng dẫn sử dụng website\nvà hệ thống Booking", size=28, bold=True,
             color=INK, align=WD_ALIGN_PARAGRAPH.CENTER, after=10, line=1.05)
    add_para(doc, "KIMMY NAILS", size=17, bold=True, color=GOLD,
             align=WD_ALIGN_PARAGRAPH.CENTER, after=18)
    add_para(doc, "Luồng khách hàng • Quản lý lịch • WhatsApp • Hủy lịch",
             size=11.5, color=MUTED, align=WD_ALIGN_PARAGRAPH.CENTER, after=66)
    add_para(doc, "Phiên bản vận hành", size=9, bold=True, color=MUTED,
             align=WD_ALIGN_PARAGRAPH.CENTER, after=2)
    add_para(doc, "22/09/2026", size=11, bold=True, color=INK,
             align=WD_ALIGN_PARAGRAPH.CENTER, after=6)
    add_para(doc, "Dành cho chủ cửa hàng và nhân viên quản lý lịch hẹn",
             size=10, italic=True, color=MUTED, align=WD_ALIGN_PARAGRAPH.CENTER, after=16)
    add_callout(doc, "Bảo mật", "Tài liệu có chứa mật khẩu quản trị. Chỉ lưu hành nội bộ và không gửi cho khách hàng.", fill="FFF3E8", accent=RED)
    doc.add_page_break()

    add_heading(doc, "1. Mục đích và phạm vi", 1)
    add_para(doc, "Tài liệu này mô tả cách website Kimmy Nails hoạt động từ lúc khách truy cập, chọn dịch vụ và đặt lịch cho đến khi dữ liệu được lưu, tin nhắn WhatsApp được tạo, lịch được hủy hoặc được nhân viên cập nhật trong trang quản trị.")
    add_list_item(doc, "Hướng dẫn khách hàng sử dụng phần Booking trên máy tính và điện thoại.", bullet_id)
    add_list_item(doc, "Giải thích quy tắc số chỗ còn lại theo từng khung giờ.", bullet_id)
    add_list_item(doc, "Hướng dẫn khách tự hủy lịch bằng liên kết bảo mật.", bullet_id)
    add_list_item(doc, "Hướng dẫn nhân viên tìm kiếm lịch sử, hủy lịch và đánh dấu hoàn thành.", bullet_id)
    add_list_item(doc, "Cung cấp cách xử lý nhanh các lỗi vận hành thường gặp.", bullet_id)

    add_heading(doc, "2. Tổng quan website", 1)
    add_para(doc, "Website là trang giới thiệu và đặt lịch trực tuyến của Kimmy Nails. Giao diện chính bằng tiếng Đức, tương thích máy tính và điện thoại. Các khu vực chính gồm:")
    add_table(
        doc,
        ["Khu vực", "Mục đích"],
        [
            ("Über uns", "Giới thiệu salon, không gian và đội ngũ."),
            ("Leistungen & Preise", "Trình bày dịch vụ và bảng giá tham khảo của salon."),
            ("Galerie", "Hiển thị hình ảnh mẫu và không gian làm việc."),
            ("Buchung", "Cho khách chọn dịch vụ, ngày, giờ và gửi yêu cầu đặt lịch."),
            ("FAQ", "Giải đáp các câu hỏi thường gặp."),
            ("Kontakt", "Hiển thị địa chỉ, số điện thoại, email và đường đi."),
        ],
        [2200, 7160],
    )
    add_callout(doc, "Lưu ý", "Giá có thể xuất hiện trong khu vực bảng giá của website, nhưng không xuất hiện trong các bước Booking, màn hình xác nhận, tin nhắn WhatsApp hoặc trang hủy lịch.")

    doc.add_page_break()
    add_heading(doc, "3. Kiến trúc và luồng dữ liệu", 1)
    add_para(doc, "Phần giao diện được hiển thị trên tên miền Kimmy Nails. Riêng dữ liệu booking được xử lý bởi API trên Vercel và lưu bền vững trong Supabase PostgreSQL.")
    flow = add_table(
        doc,
        ["1. Khách hàng", "2. Website", "3. Vercel API", "4. Supabase", "5. Kết quả"],
        [
            ("Chọn dịch vụ và thời gian", "Gửi yêu cầu booking", "Kiểm tra dữ liệu và sức chứa", "Đọc/ghi lịch hẹn", "Mở WhatsApp và tạo link hủy"),
        ],
        [1450, 1750, 1900, 1700, 2560],
        [WD_ALIGN_PARAGRAPH.CENTER] * 5,
    )
    for cell in flow.rows[1].cells:
        set_cell_shading(cell, CREAM)
    add_list_item(doc, "Website không tự lưu số chỗ trống trong trình duyệt; mỗi lần chọn hoặc tải lại ngày, dữ liệu được lấy lại từ Supabase.", bullet_id)
    add_list_item(doc, "Vercel kiểm tra lại số booking trước khi ghi dữ liệu; Supabase có cơ chế khóa và giới hạn ở mức database để tránh hai khách đặt đồng thời vượt sức chứa.", bullet_id)
    add_list_item(doc, "Booking có trạng thái cancelled vẫn được giữ lại làm lịch sử, nhưng không chiếm chỗ trong khung giờ.", bullet_id)

    doc.add_page_break()
    add_heading(doc, "4. Luồng đặt lịch dành cho khách hàng", 1)
    add_heading(doc, "Bước 1 — Mở phần Booking", 2)
    add_para(doc, "Khách bấm nút “Buchen”, “Jetzt Buchen” hoặc mục “Buchung” trên thanh điều hướng. Trang tự cuộn đến khu vực đặt lịch.")

    add_heading(doc, "Bước 2 — Chọn dịch vụ", 2)
    add_para(doc, "Khách chọn một hoặc nhiều dịch vụ. Nút đã chọn được đánh dấu. Nút “Weiter zur Buchung” chỉ hoạt động sau khi có ít nhất một dịch vụ.")
    add_para(doc, "Các nhóm dịch vụ hiện có gồm Pulver-Gel, Flüssig-Gel, Handpflege, Fußpflege, Nagelkunst và Wimpern. Trong bước Booking chỉ hiển thị tên dịch vụ, không hiển thị giá.")

    add_heading(doc, "Bước 3 — Chọn ngày", 2)
    add_para(doc, "Khách chọn ngày tại trường “Datum”. Không thể chọn ngày trong quá khứ. Ngay sau khi chọn ngày, website gọi API để lấy danh sách chỗ trống trực tiếp từ database.")

    add_heading(doc, "Bước 4 — Chọn giờ", 2)
    add_para(doc, "Trường “Uhrzeit” ban đầu hiển thị “Zuerst Datum wählen” — nghĩa là phải chọn ngày trước. Sau khi tải xong, danh sách giờ hoạt động từ 09:00 đến 19:55, cách nhau 5 phút.")
    add_table(
        doc,
        ["Loại khung giờ", "Sức chứa", "Cách hiển thị", "Khi đầy"],
        [
            ("Phút = 00\nVí dụ: 10:00", "Tối đa 3 khách", "Noch 3/2/1 Plätze\n(Còn 3/2/1 chỗ)", "Ausgebucht; không thể chọn"),
            ("Phút khác 00\nVí dụ: 10:05, 10:30", "Tối đa 1 khách", "Verfügbar\n(Còn chỗ)", "Ausgebucht; không thể chọn"),
        ],
        [1900, 1600, 3000, 2860],
        [WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT],
    )
    add_callout(doc, "Quy tắc quan trọng", "Lịch có trạng thái confirmed hoặc completed được tính vào sức chứa. Lịch cancelled không được tính; khi một lịch bị hủy, chỗ trống được mở lại tự động.")

    doc.add_page_break()
    add_heading(doc, "Bước 5 — Nhập thông tin khách", 2)
    add_table(
        doc,
        ["Nhãn trên website", "Ý nghĩa", "Bắt buộc"],
        [
            ("Name", "Tên khách hàng", "Có"),
            ("Telefon", "Số điện thoại liên hệ", "Có"),
            ("Notiz (optional)", "Ghi chú hoặc yêu cầu thêm", "Không"),
        ],
        [2500, 4860, 2000],
        [WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER],
    )

    add_heading(doc, "Bước 6 — Kiểm tra và xác nhận", 2)
    add_para(doc, "Khách bấm “Termin prüfen” để xem lại thông tin. Màn hình xác nhận chỉ gồm tên khách, số điện thoại, dịch vụ, ngày, giờ và ghi chú nếu có; không có giá hoặc tổng tiền.")
    add_para(doc, "Khi bấm “Termin bestätigen”, server kiểm tra lại database. Nếu khung giờ vừa đầy do một khách khác đặt đồng thời, booking bị từ chối và khách được yêu cầu chọn giờ khác. Nếu còn chỗ, booking được lưu với trạng thái confirmed.")

    add_heading(doc, "Bước 7 — Gửi thông tin qua WhatsApp", 2)
    add_para(doc, "Sau khi lưu thành công, website mở tin nhắn WhatsApp đã được điền sẵn tới số +49 821 44917851. Tin nhắn gồm mã booking, tên, số điện thoại, dịch vụ, ngày, giờ, ghi chú và liên kết quản lý/hủy lịch.")
    add_callout(doc, "Không trùng dữ liệu", "WhatsApp chỉ mở sau khi booking đã được lưu thành công trong Supabase. Nếu trình duyệt chặn cửa sổ mới, khách vẫn có thể bấm nút “Buchung per WhatsApp senden” trên màn hình thành công.")

    doc.add_page_break()
    add_heading(doc, "5. Luồng hủy lịch dành cho khách hàng", 1)
    add_para(doc, "Mỗi booking được tạo kèm booking_id và cancel_token riêng. Liên kết hủy có dạng /booking/cancel?id=BOOKING_ID&token=CANCEL_TOKEN và được đưa vào tin nhắn WhatsApp.")
    add_list_item(doc, "Khách mở liên kết “Termin verwalten / stornieren” hoặc link trong WhatsApp.", cancel_flow_number_ids[0])
    add_list_item(doc, "Trang hiển thị tên khách, dịch vụ, ngày và giờ để khách kiểm tra.", cancel_flow_number_ids[1])
    add_list_item(doc, "Khách bấm “Cancel Booking / Hủy lịch”.", cancel_flow_number_ids[2])
    add_list_item(doc, "Hệ thống hỏi: “Bạn có chắc chắn muốn hủy lịch này không?”.", cancel_flow_number_ids[3])
    add_list_item(doc, "Sau khi xác nhận, trạng thái được đổi thành cancelled và hiển thị thông báo hủy thành công.", cancel_flow_number_ids[4])
    add_list_item(doc, "Database giữ nguyên booking trong lịch sử; chỗ trống của khung giờ được mở lại ngay.", cancel_flow_number_ids[5])
    add_callout(doc, "Giới hạn", "Không thể hủy lại booking đã cancelled và không thể hủy booking đã completed. Link chỉ hợp lệ khi cả booking_id và cancel_token khớp.")

    add_heading(doc, "6. Quản trị Booking", 1)
    add_heading(doc, "6.1 Thông tin đăng nhập", 2)
    access = doc.add_table(rows=3, cols=2)
    values = [
        ("Trang quản trị", "https://www.kimmynail.de/admin/bookings"),
        ("Mật khẩu / Admin Key", "Aa@123456789"),
        ("Mục đích", "Tìm kiếm lịch sử, hủy lịch và đánh dấu lịch hoàn thành"),
    ]
    for row, (label, value) in zip(access.rows, values):
        set_cell_shading(row.cells[0], GOLD_LIGHT)
        p1 = row.cells[0].paragraphs[0]
        p1.paragraph_format.space_after = Pt(0)
        set_run_font(p1.add_run(label), size=10, bold=True)
        p2 = row.cells[1].paragraphs[0]
        p2.paragraph_format.space_after = Pt(0)
        if value.startswith("https://"):
            add_hyperlink(p2, value, value)
        else:
            set_run_font(p2.add_run(value), size=10, bold=("Aa@" in value), color=RED if "Aa@" in value else INK)
    set_table_geometry(access, [2600, 6760])
    set_table_borders(access)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    add_callout(doc, "Bảo mật mật khẩu", "Không gửi mật khẩu admin cho khách hàng, không đăng công khai và không lưu trên thiết bị lạ. Nếu nghi ngờ bị lộ, cần thay ADMIN_KEY trên Vercel và redeploy ngay.", fill="FFF3E8", accent=RED)

    add_heading(doc, "6.2 Cách tìm kiếm lịch sử booking", 2)
    add_list_item(doc, "Mở trang quản trị bằng đường dẫn ở trên.", admin_search_number_ids[0])
    add_list_item(doc, "Nhập Aa@123456789 vào ô “Admin-Schlüssel”.", admin_search_number_ids[1])
    add_list_item(doc, "Nếu cần, chọn “Datum” để lọc theo ngày và “Status” để lọc theo trạng thái.", admin_search_number_ids[2])
    add_list_item(doc, "Bấm “Aktualisieren”. Hệ thống gửi Admin Key đến API và tải dữ liệu mới nhất từ Supabase.", admin_search_number_ids[3])
    add_list_item(doc, "Để xem toàn bộ lịch sử, để trống Datum, chọn “Alle Status”, rồi bấm Aktualisieren.", admin_search_number_ids[4])

    add_heading(doc, "6.3 Thông tin hiển thị", 2)
    add_para(doc, "Danh sách booking gồm Name, Telefon, Leistungen, Datum, Uhrzeit, Status và Aktion. Booking bị hủy vẫn giữ trong lịch sử, được làm mờ và gắn badge Cancelled.")

    add_heading(doc, "6.4 Hủy lịch từ trang admin", 2)
    add_list_item(doc, "Tìm dòng booking cần hủy bằng bộ lọc ngày hoặc trạng thái.", admin_cancel_number_ids[0])
    add_list_item(doc, "Tại cột Aktion, bấm “Cancel Booking”.", admin_cancel_number_ids[1])
    add_list_item(doc, "Đọc popup xác nhận và bấm đồng ý nếu đúng booking.", admin_cancel_number_ids[2])
    add_list_item(doc, "Hệ thống đổi trạng thái thành cancelled, giữ nguyên lịch sử và mở lại chỗ trống.", admin_cancel_number_ids[3])
    add_callout(doc, "Kiểm tra trước khi hủy", "Đối chiếu tối thiểu tên khách, số điện thoại, dịch vụ, ngày và giờ. Không thể hoàn tác trực tiếp bằng giao diện sau khi đã hủy.", fill="FFF3E8", accent=RED)

    add_heading(doc, "6.5 Đánh dấu hoàn thành", 2)
    add_para(doc, "Với booking đang confirmed, nhân viên có thể bấm “Completed” sau khi khách đã sử dụng dịch vụ. Booking completed vẫn chiếm sức chứa của khung giờ đó và được giữ trong lịch sử.")

    doc.add_page_break()
    add_heading(doc, "7. Ý nghĩa trạng thái booking", 1)
    add_table(
        doc,
        ["Trạng thái", "Ý nghĩa", "Tính vào sức chứa", "Thao tác tiếp theo"],
        [
            ("Confirmed", "Đã đặt thành công, đang chờ phục vụ", "Có", "Admin có thể Cancel hoặc Completed"),
            ("Cancelled", "Đã hủy bởi khách hoặc admin", "Không", "Giữ trong lịch sử; không hủy lại"),
            ("Completed", "Khách đã hoàn tất dịch vụ", "Có", "Không thể hủy từ giao diện khách"),
        ],
        [1800, 3000, 2000, 2560],
        [WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT],
    )

    add_heading(doc, "8. Kiểm tra an toàn dữ liệu và sức chứa", 1)
    add_list_item(doc, "Frontend tải availability để disable giờ đầy và hiển thị số chỗ còn lại.", bullet_id)
    add_list_item(doc, "Backend kiểm tra lại booking_date + booking_time ngay khi khách xác nhận.", bullet_id)
    add_list_item(doc, "Database áp dụng khóa giao dịch theo ngày và giờ để tránh vượt sức chứa khi có nhiều yêu cầu đồng thời.", bullet_id)
    add_list_item(doc, "Phút bằng 00 có tối đa 3 booking; phút khác 00 có tối đa 1 booking.", bullet_id)
    add_list_item(doc, "Chỉ booking cancelled được loại khỏi phép đếm sức chứa.", bullet_id)
    add_list_item(doc, "Không xóa vật lý booking khi hủy; toàn bộ lịch sử vẫn nằm trong Supabase.", bullet_id)

    add_heading(doc, "9. Xử lý sự cố thường gặp", 1)
    add_table(
        doc,
        ["Hiện tượng", "Nguyên nhân thường gặp", "Cách xử lý"],
        [
            ("Không tải được danh sách giờ", "Trang đang dùng bản cache cũ hoặc API tạm thời không phản hồi", "Ctrl + F5; publish lại Webcake; kiểm tra Vercel deployment và DATABASE_URL"),
            ("Unexpected token '<'", "Trình duyệt nhận HTML thay vì JSON từ API", "Đảm bảo booking gọi API Vercel; làm mới cache của trang/iframe"),
            ("Keine Zeiten verfügbar", "Không tải được availability hoặc cấu hình database lỗi", "Kiểm tra Supabase, biến DATABASE_URL và log Vercel"),
            ("Admin báo Unauthorized", "Thiếu hoặc sai Admin Key", "Nhập đúng Aa@123456789 rồi bấm Aktualisieren"),
            ("Không mở được WhatsApp", "Trình duyệt chặn popup", "Bấm nút Buchung per WhatsApp senden trên màn hình thành công"),
            ("Khung giờ báo đầy", "Đã đạt capacity hoặc có khách vừa đặt đồng thời", "Chọn giờ khác; kiểm tra lịch confirmed/completed trong admin"),
            ("Hủy lịch nhưng chưa thấy chỗ trống", "Trang availability chưa tải lại", "Reload trang hoặc chọn lại ngày để lấy dữ liệu mới từ database"),
        ],
        [2250, 3300, 3810],
    )

    add_heading(doc, "10. Checklist vận hành hằng ngày", 1)
    add_list_item(doc, "Mở trang admin và nhập Admin Key.", bullet_id)
    add_list_item(doc, "Bấm Aktualisieren để xem các booking mới nhất.", bullet_id)
    add_list_item(doc, "Đối chiếu tin nhắn WhatsApp với booking đã lưu trong danh sách.", bullet_id)
    add_list_item(doc, "Cập nhật Completed sau khi khách hoàn tất dịch vụ.", bullet_id)
    add_list_item(doc, "Khi khách báo hủy, dùng Cancel Booking thay vì xóa dữ liệu.", bullet_id)
    add_list_item(doc, "Không chia sẻ mật khẩu admin hoặc liên kết hủy riêng của khách.", bullet_id)
    add_list_item(doc, "Nếu website báo lỗi, ghi lại thời gian, ngày/giờ booking và ảnh màn hình để kiểm tra log nhanh hơn.", bullet_id)

    add_heading(doc, "11. Thông tin tham chiếu nhanh", 1)
    ref = doc.add_table(rows=4, cols=2)
    ref_values = [
        ("Website", "https://www.kimmynail.de/"),
        ("Booking", "https://www.kimmynail.de/#booking"),
        ("Admin Booking", "https://www.kimmynail.de/admin/bookings"),
        ("WhatsApp", "+49 821 44917851"),
    ]
    for row, (label, value) in zip(ref.rows, ref_values):
        set_cell_shading(row.cells[0], GOLD_LIGHT)
        p1 = row.cells[0].paragraphs[0]
        p1.paragraph_format.space_after = Pt(0)
        set_run_font(p1.add_run(label), size=10, bold=True)
        p2 = row.cells[1].paragraphs[0]
        p2.paragraph_format.space_after = Pt(0)
        if value.startswith("https://"):
            add_hyperlink(p2, value, value)
        else:
            set_run_font(p2.add_run(value), size=10)
    set_table_geometry(ref, [2600, 6760])
    set_table_borders(ref)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)

    add_callout(doc, "Kết luận", "Luồng chuẩn là: khách chọn dịch vụ → chọn ngày/giờ → nhập thông tin → xác nhận → booking được lưu vào Supabase → WhatsApp mở với nội dung sẵn → khách hoặc admin có thể hủy mà không xóa lịch sử.", fill=GOLD_LIGHT, accent=GOLD)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()
